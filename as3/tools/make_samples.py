"""Create two clearly-fake sample workbooks that follow the A2 data contract (7 headers).

All names, policy numbers and hospitals are invented placeholders, not real people.
Run with a Python that has openpyxl:  python3 tools/make_samples.py
"""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook

HEADERS = ["customer_name", "policy_number", "coverage_start_date", "coverage_end_date",
           "hospital_name", "claim_status", "document_request"]
D = datetime
OUT = Path(__file__).resolve().parent.parent / "samples"
OUT.mkdir(exist_ok=True)

ok_rows = [
    ["คุณตัวอย่าง หนึ่ง", "POL-TEST-0001", D(2026, 1, 1), D(2026, 12, 31), "โรงพยาบาลตัวอย่างเอ", "REQUEST_DOC",
     '["สำเนาใบเสร็จค่ารักษาพยาบาล","ใบรับรองแพทย์"]'],
    ["คุณตัวอย่าง สอง", "POL-TEST-0002", D(2026, 3, 15), D(2027, 3, 14), "โรงพยาบาลตัวอย่างบี", "APPROVED", "[]"],
    ["คุณตัวอย่าง สาม", "POL-TEST-0003", D(2026, 5, 10), D(2027, 5, 9), "โรงพยาบาลตัวอย่างซี", "REQUEST_DOC",
     '["รายงานการรักษาจากแพทย์","สำเนาประวัติการรักษา","ผลตรวจทางห้องปฏิบัติการ"]'],
    ["คุณตัวอย่าง สี่", "POL-TEST-0004", D(2026, 6, 1), D(2027, 5, 31), "โรงพยาบาลตัวอย่างดี", "REJECTED", "[]"],
    ["คุณตัวอย่าง ห้า", "POL-TEST-0005", D(2026, 8, 1), D(2027, 7, 31), "โรงพยาบาลตัวอย่างอี", "REQUEST_DOC",
     '["ใบรับรองแพทย์"]'],
]
bad_rows = [
    ["คุณตัวอย่าง หก", "POL-TEST-0006", D(2027, 1, 1), D(2026, 1, 1), "โรงพยาบาลตัวอย่างเอ", "REQUEST_DOC", '["ใบรับรองแพทย์"]'],
    ["คุณตัวอย่าง เจ็ด", "POL-TEST-0007", D(2026, 1, 1), D(2026, 12, 31), "โรงพยาบาลตัวอย่างบี", "REQUEST_DOCS", '["ใบรับรองแพทย์"]'],
    ["คุณตัวอย่าง แปด", "POL-TEST-0008", D(2026, 1, 1), D(2026, 12, 31), "โรงพยาบาลตัวอย่างซี", "REQUEST_DOC", "[]"],
    ["", "POL-TEST-0009", D(2026, 1, 1), D(2026, 12, 31), "โรงพยาบาลตัวอย่างดี", "REQUEST_DOC", '["ใบรับรองแพทย์"]'],
    ["คุณตัวอย่าง สิบ", "POL-TEST-0010", D(2026, 1, 1), D(2026, 12, 31), "โรงพยาบาลตัวอย่างอี", "REQUEST_DOC", "ใบรับรองแพทย์"],
    ["คุณตัวอย่าง สิบเอ็ด", "POL-TEST-0011", "1 มกราคม 2569", D(2026, 12, 31), "โรงพยาบาลตัวอย่างเอฟ", "REQUEST_DOC", '["ใบรับรองแพทย์"]'],
    ["คุณตัวอย่าง สิบสอง", "POL-TEST-0012", D(2026, 1, 1), D(2026, 12, 31), "โรงพยาบาลตัวอย่างจี", "REQUEST_DOC", '["ใบรับรองแพทย์"]'],
]

def write(name: str, rows: list[list]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    for col in ("C", "D"):
        for cell in ws[col][1:]:
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd"
    wb.save(OUT / name)
    print("wrote", OUT / name)

write("sample_request_ok.xlsx", ok_rows)
write("sample_request_with_mistakes.xlsx", bad_rows)
