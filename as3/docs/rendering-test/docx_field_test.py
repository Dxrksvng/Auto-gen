"""DOCX field-substitution test (NOT a PDF rendering test).

Question: does a Thai field name written as {{ชื่อลูกค้า}} survive Word-style
run splitting and get replaced by docxtpl/Jinja? Also: what happens with a
misspelled field, and does a name with a space or tone mark work?

Run: uv run --with python-docx --with docxtpl python docx_field_test.py
"""
import json
from pathlib import Path
from docx import Document
from docxtpl import DocxTemplate
from jinja2 import Undefined, StrictUndefined
import jinja2

OUT = Path(__file__).parent / "out"
OUT.mkdir(exist_ok=True)
CTX = {"ชื่อลูกค้า": "คุณตัวอย่าง หนึ่ง", "เลขกรมธรรม์": "POL-DEMO-0001", "ชื่อโรงพยาบาล": "ทดสอบ"}

def make(path, paragraphs):
    d = Document()
    for runs in paragraphs:
        p = d.add_paragraph()
        for r in runs:
            p.add_run(r)
    d.save(path)

def render(name, paragraphs, strict=False):
    src, dst = OUT / f"{name}_in.docx", OUT / f"{name}_out.docx"
    make(src, paragraphs)
    t = DocxTemplate(src)
    try:
        env = jinja2.Environment(undefined=StrictUndefined) if strict else None
        t.render(CTX, jinja_env=env) if env else t.render(CTX)
        t.save(dst)
        return [p.text for p in Document(dst).paragraphs], None
    except Exception as e:  # report exactly what the author would hit
        return None, f"{type(e).__name__}: {str(e)[:160]}"

CASES = {
    "A_one_run": [["เรียน {{ชื่อลูกค้า}}"]],
    "B_split_runs_like_word": [["เรียน {{ชื่อ", "ลูกค้า}}"]],
    "C_split_inside_braces": [["เรียน {", "{ชื่อลูกค้า}", "}"]],
    "D_typo_field": [["เรียน {{ชื่อลูกค้าา}}"]],
    "E_field_with_space": [["เรียน {{ชื่อ ลูกค้า}}"]],
    "F_unknown_logic": [["{% if ชื่อลูกค้า %}ถึงลูกค้า{% endif %}"]],
}
results = {}
for k, v in CASES.items():
    out, err = render(k, v)
    results[k] = {"output": out, "error": err}
out, err = render("D_typo_field_strict", CASES["D_typo_field"], strict=True)
results["D_typo_field_strict"] = {"output": out, "error": err}
(Path(__file__).parent / "docx_field_results.json").write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps(results, ensure_ascii=False, indent=2))
