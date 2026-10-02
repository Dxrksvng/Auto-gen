"""A4: document names come from an approved mapping; unknown names are never printed raw."""

import re
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from normalize import DEFAULT_DOCUMENT_NAMES_PATH, UnknownDocumentError, load_document_names

THAI = re.compile(r"[฀-๿]")
SAMPLE_XLSX = Path(__file__).resolve().parents[2] / "data" / "Example claim table.xlsx"


@pytest.fixture(scope="module")
def names():
    return load_document_names()


def test_default_config_file_exists():
    assert DEFAULT_DOCUMENT_NAMES_PATH.exists()
    assert DEFAULT_DOCUMENT_NAMES_PATH.as_posix().endswith("config/document_names.yml")


def test_english_name_is_mapped_to_approved_thai(names):
    assert names.resolve("Medical Report") == "รายงานทางการแพทย์"


def test_matching_ignores_case_and_extra_whitespace(names):
    assert names.resolve("  medical   REPORT ") == "รายงานทางการแพทย์"


def test_approved_thai_name_maps_to_itself(names):
    assert names.resolve("ใบรับรองแพทย์") == "ใบรับรองแพทย์"


def test_unknown_name_raises_with_the_original_name(names):
    with pytest.raises(UnknownDocumentError) as info:
        names.resolve("Discharge Summary")
    assert info.value.name == "Discharge Summary"


def test_every_approved_wording_is_thai():
    raw = yaml.safe_load(DEFAULT_DOCUMENT_NAMES_PATH.read_text(encoding="utf-8"))
    for source, approved in raw["names"].items():
        assert THAI.search(approved), f"approved wording for {source!r} has no Thai text"


def test_all_document_names_in_the_sample_workbook_are_known(names):
    import json

    sheet = load_workbook(SAMPLE_XLSX, data_only=True)["Sheet1"]
    found = set()
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if row[6]:
            found.update(json.loads(row[6]))
    assert found, "sample workbook should contain document names"
    for name in found:
        names.resolve(name)
