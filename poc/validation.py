import json
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from models import LetterData, RowIssue


HEADERS = (
    "customer_name",
    "policy_number",
    "coverage_start_date",
    "coverage_end_date",
    "hospital_name",
    "claim_status",
    "document_request",
)
OBSERVED_STATUSES = {"REQUEST_DOC", "APPROVED", "REJECTED"}


class FileValidationError(ValueError):
    pass


def load_and_validate(path: Path) -> tuple[list[LetterData], list[RowIssue], list[dict]]:
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise FileValidationError(f"Workbook cannot be read: {exc}") from exc

    if "Sheet1" not in workbook.sheetnames:
        raise FileValidationError("Expected worksheet 'Sheet1' was not found")
    sheet = workbook["Sheet1"]
    rows = sheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration as exc:
        raise FileValidationError("Workbook is empty") from exc

    headers = tuple(value for value in raw_headers if value is not None)
    if headers != HEADERS:
        raise FileValidationError(f"Headers must exactly match: {', '.join(HEADERS)}")

    letters: list[LetterData] = []
    issues: list[RowIssue] = []
    skipped: list[dict] = []
    for source_row, values in enumerate(rows, start=2):
        values = values[: len(HEADERS)]
        if not any(value is not None and value != "" for value in values):
            continue
        row = dict(zip(HEADERS, values))
        row_issues = validate_row(source_row, row)
        if row_issues:
            issues.extend(row_issues)
            continue
        if row["claim_status"] != "REQUEST_DOC":
            skipped.append({"source_row": source_row, "reason": f"status={row['claim_status']}"})
            continue
        documents = tuple(item.strip() for item in json.loads(row["document_request"]))
        letters.append(
            LetterData(
                source_row=source_row,
                customer_name=row["customer_name"].strip(),
                policy_number=row["policy_number"].strip(),
                coverage_start_date=row["coverage_start_date"],
                coverage_end_date=row["coverage_end_date"],
                hospital_name=row["hospital_name"].strip(),
                requested_documents=documents,
            )
        )
    return letters, issues, skipped


def validate_row(source_row: int, row: dict) -> list[RowIssue]:
    issues: list[RowIssue] = []
    for field in ("customer_name", "policy_number", "hospital_name", "claim_status"):
        if not isinstance(row[field], str) or not row[field].strip():
            issues.append(RowIssue(source_row, field, "REQUIRED", "Value must be a non-empty string"))
    for field in ("coverage_start_date", "coverage_end_date"):
        if not isinstance(row[field], datetime):
            issues.append(RowIssue(source_row, field, "INVALID_DATE", "Value must be an Excel date"))
    if isinstance(row["coverage_start_date"], datetime) and isinstance(row["coverage_end_date"], datetime):
        if row["coverage_start_date"] > row["coverage_end_date"]:
            issues.append(RowIssue(source_row, "coverage_end_date", "DATE_ORDER", "End date precedes start date"))
    if row["claim_status"] not in OBSERVED_STATUSES:
        issues.append(RowIssue(source_row, "claim_status", "UNKNOWN_STATUS", "Status needs business-owner review"))
    try:
        documents = json.loads(row["document_request"])
        valid_list = isinstance(documents, list) and all(isinstance(x, str) and x.strip() for x in documents)
    except (TypeError, json.JSONDecodeError):
        valid_list = False
        documents = None
    if not valid_list:
        issues.append(RowIssue(source_row, "document_request", "INVALID_LIST", "Must be a JSON list of non-empty strings"))
    elif row["claim_status"] == "REQUEST_DOC" and not documents:
        issues.append(RowIssue(source_row, "document_request", "REQUIRED_FOR_REQUEST", "REQUEST_DOC needs at least one document"))
    return issues

